import os
import re
import time
import uuid
from io import BytesIO

import pytest
import requests
from dotenv import dotenv_values
from openpyxl import load_workbook


# Auth + environment helpers
BASE_URL = (os.environ.get("REACT_APP_BACKEND_URL") or dotenv_values("/app/frontend/.env").get("REACT_APP_BACKEND_URL", "")).rstrip("/")
assert BASE_URL, "REACT_APP_BACKEND_URL is required"


def _read_admin_credentials():
    text = open("/app/memory/test_credentials.md", "r", encoding="utf-8").read()
    email = re.search(r"-\s*Email:\s*`([^`]+)`", text)
    password = re.search(r"-\s*Password:\s*`([^`]+)`", text)
    if not email or not password:
        pytest.skip("Admin credentials missing from /app/memory/test_credentials.md")
    return email.group(1), password.group(1)


def _create_product_payload(label: str, stock: int = 3, price: int = 15000):
    return {
        "title": f"TEST_{label}_{uuid.uuid4().hex[:6]}",
        "brand": "Lenovo",
        "category": "Business",
        "price": price,
        "compare_at_price": price + 2000,
        "condition_grade": "A",
        "condition_description": "TEST stock-safe description",
        "processor": "Intel Core i5",
        "ram_gb": 8,
        "storage_type": "SSD",
        "storage_gb": 256,
        "display": "14-inch Full HD",
        "gpu": "Integrated",
        "battery_health": "Tested",
        "operating_system": "Windows 11 Pro",
        "warranty_months": 3,
        "stock_quantity": stock,
        "image_url": "https://images.unsplash.com/photo-1496181133206-80ce9b88a853?auto=format&fit=crop&w=1200&q=80",
        "status": "active",
    }


def _find_order(session, headers, order_id: str):
    response = session.get(f"{BASE_URL}/api/admin/orders", headers=headers)
    assert response.status_code == 200, response.text
    for order in response.json():
        if order["order_id"] == order_id:
            return order
    return None


def _wait_for_notification(session, headers, order_id: str, recipient: str, timeout_seconds: int = 20):
    deadline = time.time() + timeout_seconds
    while time.time() < deadline:
        order = _find_order(session, headers, order_id)
        if order:
            state = order.get("notifications", {}).get(recipient, {}).get("status")
            if state in {"accepted", "not_configured", "failed", "unknown"}:
                return order
        time.sleep(1)
    return _find_order(session, headers, order_id)


def _wait_for_customer_final(session, headers, order_id: str, timeout_seconds: int = 20):
    deadline = time.time() + timeout_seconds
    while time.time() < deadline:
        order = _find_order(session, headers, order_id)
        if order:
            state = order.get("notifications", {}).get("customer", {}).get("status")
            if state in {"accepted", "unknown", "failed"}:
                return order
        time.sleep(1)
    return _find_order(session, headers, order_id)


@pytest.fixture(scope="module")
def client():
    session = requests.Session()
    session.headers.update({"Content-Type": "application/json"})
    return session


@pytest.fixture(scope="module")
def admin_creds():
    return _read_admin_credentials()


@pytest.fixture(scope="module")
def auth_headers(client, admin_creds):
    email, password = admin_creds
    response = client.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["email"] == email
    assert data["role"] == "admin"
    return {"Authorization": f"Bearer {data['token']}"}


@pytest.fixture(scope="module")
def test_state(client, auth_headers):
    # Shared module state for safe cleanup/revert of settings and temporary products.
    settings = client.get(f"{BASE_URL}/api/settings")
    assert settings.status_code == 200, settings.text
    original_settings = settings.json()
    state = {"original_settings": original_settings, "created_product_ids": []}
    yield state
    for product_id in state["created_product_ids"]:
        client.delete(f"{BASE_URL}/api/admin/products/{product_id}", headers=auth_headers)
    restore_payload = {
        "store_name": original_settings.get("store_name", "Bhose Technology"),
        "owner_email": original_settings.get("owner_email"),
        "phone": original_settings.get("phone", ""),
        "address": original_settings.get("address", ""),
        "shipping_policy": original_settings.get("shipping_policy", ""),
        "return_policy": original_settings.get("return_policy", ""),
        "warranty_default": original_settings.get("warranty_default", ""),
        "hero_title": original_settings.get("hero_title", ""),
        "hero_subtitle": original_settings.get("hero_subtitle", ""),
    }
    client.put(f"{BASE_URL}/api/admin/settings", headers=auth_headers, json=restore_payload)


# Auth + protection checks
def test_auth_login_me_and_admin_guard(client, auth_headers, admin_creds):
    me = client.get(f"{BASE_URL}/api/auth/me", headers=auth_headers)
    assert me.status_code == 200, me.text
    assert me.json() == {"email": admin_creds[0], "role": "admin"}
    protected = client.get(f"{BASE_URL}/api/admin/products")
    assert protected.status_code == 401


def test_bruteforce_lockout_after_failed_attempts(client):
    identifier = f"lock-{uuid.uuid4().hex[:6]}@example.com"
    statuses = []
    for _ in range(6):
        response = client.post(f"{BASE_URL}/api/auth/login", json={"email": identifier, "password": "WrongPass123!"})
        statuses.append(response.status_code)
    assert statuses[:5] == [401, 401, 401, 401, 401]
    assert statuses[5] == 429


# Product + stock + soft-delete checks
def test_product_edit_requires_expected_stock_and_persists(client, auth_headers, test_state):
    created = client.post(f"{BASE_URL}/api/admin/products", headers=auth_headers, json=_create_product_payload("edit", stock=10))
    assert created.status_code == 200, created.text
    product = created.json()
    test_state["created_product_ids"].append(product["product_id"])
    update_payload = {**product, "title": product["title"] + "_UPDATED", "expected_stock": product["stock_quantity"]}
    updated = client.put(f"{BASE_URL}/api/admin/products/{product['product_id']}", headers=auth_headers, json=update_payload)
    assert updated.status_code == 200, updated.text
    assert updated.json()["title"].endswith("_UPDATED")


def test_patch_stock_validation_and_same_count_allowed(client, auth_headers, test_state):
    created = client.post(f"{BASE_URL}/api/admin/products", headers=auth_headers, json=_create_product_payload("stock", stock=4))
    assert created.status_code == 200, created.text
    product = created.json()
    test_state["created_product_ids"].append(product["product_id"])

    bad_negative = client.patch(
        f"{BASE_URL}/api/admin/products/{product['product_id']}/stock",
        headers=auth_headers,
        json={"stock_quantity": -1, "expected_stock": product["stock_quantity"]},
    )
    assert bad_negative.status_code == 422

    bad_fraction = client.patch(
        f"{BASE_URL}/api/admin/products/{product['product_id']}/stock",
        headers=auth_headers,
        json={"stock_quantity": 1.5, "expected_stock": product["stock_quantity"]},
    )
    assert bad_fraction.status_code == 422

    same_count = client.patch(
        f"{BASE_URL}/api/admin/products/{product['product_id']}/stock",
        headers=auth_headers,
        json={"stock_quantity": product["stock_quantity"], "expected_stock": product["stock_quantity"]},
    )
    assert same_count.status_code == 200, same_count.text
    assert same_count.json()["stock_quantity"] == product["stock_quantity"]


def test_stale_stock_conflict_does_not_overwrite_latest_stock(client, auth_headers, test_state):
    created = client.post(f"{BASE_URL}/api/admin/products", headers=auth_headers, json=_create_product_payload("conflict", stock=10))
    assert created.status_code == 200, created.text
    product = created.json()
    test_state["created_product_ids"].append(product["product_id"])

    fresh_stock = client.patch(
        f"{BASE_URL}/api/admin/products/{product['product_id']}/stock",
        headers=auth_headers,
        json={"stock_quantity": 8, "expected_stock": 10},
    )
    assert fresh_stock.status_code == 200, fresh_stock.text

    stale_payload = {**product, "title": product["title"] + "_STALE_SAVE", "expected_stock": 10}
    stale_once = client.put(f"{BASE_URL}/api/admin/products/{product['product_id']}", headers=auth_headers, json=stale_payload)
    stale_twice = client.put(f"{BASE_URL}/api/admin/products/{product['product_id']}", headers=auth_headers, json=stale_payload)
    assert stale_once.status_code == 409
    assert stale_twice.status_code == 409

    inventory = client.get(f"{BASE_URL}/api/admin/products", headers=auth_headers)
    assert inventory.status_code == 200
    latest = next(item for item in inventory.json() if item["product_id"] == product["product_id"])
    assert latest["stock_quantity"] == 8
    assert latest["title"] == product["title"]


def test_unknown_product_returns_404_for_edit_stock_and_delete(client, auth_headers):
    unknown = f"missing-{uuid.uuid4().hex}"
    put_resp = client.put(f"{BASE_URL}/api/admin/products/{unknown}", headers=auth_headers, json={**_create_product_payload("missing"), "expected_stock": 0})
    patch_resp = client.patch(f"{BASE_URL}/api/admin/products/{unknown}/stock", headers=auth_headers, json={"stock_quantity": 1, "expected_stock": 0})
    delete_resp = client.delete(f"{BASE_URL}/api/admin/products/{unknown}", headers=auth_headers)
    assert put_resp.status_code == 404
    assert patch_resp.status_code == 404
    assert delete_resp.status_code == 404


def test_soft_delete_removes_from_catalog_and_detail(client, auth_headers):
    created = client.post(f"{BASE_URL}/api/admin/products", headers=auth_headers, json=_create_product_payload("delete", stock=1))
    assert created.status_code == 200, created.text
    product = created.json()

    removed = client.delete(f"{BASE_URL}/api/admin/products/{product['product_id']}", headers=auth_headers)
    assert removed.status_code == 200, removed.text
    assert removed.json()["status"] == "deleted"

    public_list = client.get(f"{BASE_URL}/api/products")
    assert public_list.status_code == 200
    assert all(item["product_id"] != product["product_id"] for item in public_list.json())
    public_detail = client.get(f"{BASE_URL}/api/products/{product['product_id']}")
    assert public_detail.status_code == 404


# Order + email notifications checks
def test_order_quantity_total_stock_and_owner_missing_status(client, auth_headers, test_state):
    created = client.post(f"{BASE_URL}/api/admin/products", headers=auth_headers, json=_create_product_payload("order", stock=5, price=11111))
    assert created.status_code == 200, created.text
    product = created.json()
    test_state["created_product_ids"].append(product["product_id"])

    order_payload = {
        "customer_name": "TEST Buyer",
        "customer_email": "delivered@resend.dev",
        "customer_phone": "9999988888",
        "address": "TEST Street 42, Business Park",
        "city": "Pune",
        "pincode": "411001",
        "items": [{"product_id": product["product_id"], "quantity": 2}],
    }
    order_resp = client.post(f"{BASE_URL}/api/orders", json=order_payload)
    assert order_resp.status_code == 200, order_resp.text
    order = order_resp.json()
    assert order["total"] == 22222
    assert order["items"][0]["quantity"] == 2

    inventory = client.get(f"{BASE_URL}/api/admin/products", headers=auth_headers)
    assert inventory.status_code == 200
    updated = next(item for item in inventory.json() if item["product_id"] == product["product_id"])
    assert updated["stock_quantity"] == 3

    settled_owner = _wait_for_notification(client, auth_headers, order["order_id"], "owner", timeout_seconds=25)
    assert settled_owner is not None
    assert settled_owner.get("notifications", {}).get("owner", {}).get("status") == "not_configured"
    settled_customer = _wait_for_customer_final(client, auth_headers, order["order_id"], timeout_seconds=25)
    assert settled_customer.get("notifications", {}).get("customer", {}).get("status") in {"accepted", "unknown", "failed"}


def test_retry_owner_email_after_temp_config_and_skip_duplicate(client, auth_headers, test_state):
    current = client.get(f"{BASE_URL}/api/settings").json()
    payload = {
        "store_name": current.get("store_name", "Bhose Technology"),
        "owner_email": "delivered@resend.dev",
        "phone": current.get("phone", ""),
        "address": current.get("address", ""),
        "shipping_policy": current.get("shipping_policy", ""),
        "return_policy": current.get("return_policy", ""),
        "warranty_default": current.get("warranty_default", ""),
        "hero_title": current.get("hero_title", ""),
        "hero_subtitle": current.get("hero_subtitle", ""),
    }
    save = client.put(f"{BASE_URL}/api/admin/settings", headers=auth_headers, json=payload)
    assert save.status_code == 200, save.text
    assert save.json()["owner_email"] == "delivered@resend.dev"

    created = client.post(f"{BASE_URL}/api/admin/products", headers=auth_headers, json=_create_product_payload("retry", stock=2, price=12000))
    assert created.status_code == 200, created.text
    product = created.json()
    test_state["created_product_ids"].append(product["product_id"])

    order_resp = client.post(
        f"{BASE_URL}/api/orders",
        json={
            "customer_name": "TEST Retry",
            "customer_email": "delivered@resend.dev",
            "customer_phone": "9999977777",
            "address": "TEST retry lane 88",
            "city": "Mumbai",
            "pincode": "400001",
            "items": [{"product_id": product["product_id"], "quantity": 1}],
        },
    )
    assert order_resp.status_code == 200, order_resp.text
    order_id = order_resp.json()["order_id"]

    settled = _wait_for_notification(client, auth_headers, order_id, "owner", timeout_seconds=25)
    assert settled is not None
    if settled.get("notifications", {}).get("owner", {}).get("status") != "accepted":
        retry = client.post(
            f"{BASE_URL}/api/admin/orders/{order_id}/retry-email",
            headers=auth_headers,
            params={"recipient": "owner"},
        )
        assert retry.status_code == 200, retry.text
        settled = _wait_for_notification(client, auth_headers, order_id, "owner", timeout_seconds=25)

    owner_result = settled.get("notifications", {}).get("owner", {})
    assert owner_result.get("status") == "accepted"
    assert owner_result.get("email_id")

    duplicate_retry = client.post(
        f"{BASE_URL}/api/admin/orders/{order_id}/retry-email",
        headers=auth_headers,
        params={"recipient": "owner"},
    )
    assert duplicate_retry.status_code == 200
    assert duplicate_retry.json().get("notifications", {}).get("owner", {}).get("status") == "accepted"


def test_test_email_rate_limit_and_documentation_download(client, auth_headers):
    first = client.post(f"{BASE_URL}/api/admin/email/test", headers=auth_headers)
    if first.status_code == 200:
        assert first.json().get("status") in {"accepted", "unknown", "failed"}
        second = client.post(f"{BASE_URL}/api/admin/email/test", headers=auth_headers)
        assert second.status_code == 429
    else:
        assert first.status_code == 429

    docs = client.get(f"{BASE_URL}/api/admin/documentation", headers=auth_headers)
    assert docs.status_code == 200, docs.text
    assert "text/markdown" in docs.headers.get("content-type", "")
    assert len(docs.text.strip()) > 200


# Reports + settings checks
def test_reports_export_mime_filename_parseable_and_unauthorized(client, auth_headers):
    unauth = client.get(f"{BASE_URL}/api/admin/orders/export", params={"format": "xlsx"})
    assert unauth.status_code == 401

    xlsx = client.get(f"{BASE_URL}/api/admin/orders/export", headers=auth_headers, params={"format": "xlsx"})
    assert xlsx.status_code == 200, xlsx.text
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in xlsx.headers.get("content-type", "")
    assert "orders-" in xlsx.headers.get("content-disposition", "")
    wb = load_workbook(BytesIO(xlsx.content))
    assert set(wb.sheetnames) >= {"Summary", "Orders", "Items"}

    pdf = client.get(f"{BASE_URL}/api/admin/orders/export", headers=auth_headers, params={"format": "pdf", "date": "2099-01-01"})
    assert pdf.status_code == 200, pdf.text
    assert "application/pdf" in pdf.headers.get("content-type", "")
    assert pdf.content.startswith(b"%PDF")
    assert len(pdf.content) > 500


def test_store_name_update_persists_and_owner_email_blank_normalizes(client, auth_headers):
    before = client.get(f"{BASE_URL}/api/settings")
    assert before.status_code == 200
    baseline = before.json()
    renamed = f"TEST_STORE_{uuid.uuid4().hex[:6]}"

    payload = {
        "store_name": renamed,
        "owner_email": "   ",
        "phone": baseline.get("phone", ""),
        "address": baseline.get("address", ""),
        "shipping_policy": baseline.get("shipping_policy", ""),
        "return_policy": baseline.get("return_policy", ""),
        "warranty_default": baseline.get("warranty_default", ""),
        "hero_title": baseline.get("hero_title", ""),
        "hero_subtitle": baseline.get("hero_subtitle", ""),
    }
    saved = client.put(f"{BASE_URL}/api/admin/settings", headers=auth_headers, json=payload)
    assert saved.status_code == 200, saved.text
    assert saved.json()["store_name"] == renamed
    assert saved.json()["owner_email"] is None

    fetched = client.get(f"{BASE_URL}/api/settings")
    assert fetched.status_code == 200
    assert fetched.json()["store_name"] == renamed

    invalid_email_payload = {**payload, "store_name": renamed, "owner_email": "not-an-email"}
    invalid = client.put(f"{BASE_URL}/api/admin/settings", headers=auth_headers, json=invalid_email_payload)
    assert invalid.status_code == 422
