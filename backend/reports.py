from datetime import datetime, time, timedelta, timezone
from html import escape
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether

STORE_TZ = ZoneInfo("Asia/Kolkata")

def day_bounds(day):
    start = datetime.combine(day, time.min, tzinfo=STORE_TZ)
    return start.astimezone(timezone.utc).isoformat(), (start + timedelta(days=1)).astimezone(timezone.utc).isoformat()

def local_date(value):
    return datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(STORE_TZ).strftime("%d %b %Y, %H:%M IST")

def excel_text(value):
    value = ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value

def workbook_bytes(orders, store_name, day):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append([excel_text(store_name), day.isoformat(), "Asia/Kolkata"])
    summary.append(["Orders", len(orders)])
    summary.append(["All order value (INR)", sum(o["total"] for o in orders)])
    summary.append(["Non-cancelled value (INR)", sum(o["total"] for o in orders if o["status"] != "CANCELLED")])
    sheet = wb.create_sheet("Orders")
    sheet.append(["Order number", "Placed (IST)", "Customer", "Email", "Phone", "Address", "City", "PIN code", "Status", "Total (INR)", "Owner email status", "Customer email status"])
    lines = wb.create_sheet("Items")
    lines.append(["Order number", "SKU", "Product", "Quantity", "Unit price (INR)", "Line total (INR)"])
    for order in orders:
        sheet.append([excel_text(order[k]) for k in ["order_number"]] + [local_date(order["created_at"])] +
            [excel_text(order[k]) for k in ["customer_name", "customer_email", "customer_phone", "address", "city", "pincode", "status"]] +
            [order["total"]] + [excel_text(order.get("notifications", {}).get(target, {}).get("status", "not_tracked")) for target in ["owner", "customer"]])
        for item in order["items"]:
            lines.append([excel_text(order["order_number"]), excel_text(item["sku"]), excel_text(item["title"]), item["quantity"], item["price"], item["quantity"] * item["price"]])
    for ws in wb:
        ws.freeze_panes = "A2"
        if ws != summary:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="087F76")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(45, max(18, max(len(str(c.value or "")) for c in col) + 2))
            for cell in col:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()

def pdf_bytes(orders, store_name, day):
    output = BytesIO()
    styles = getSampleStyleSheet()
    styles["BodyText"].fontSize = 9
    styles["BodyText"].leading = 13
    def p(value):
        return Paragraph(escape(str(value)), styles["BodyText"])
    story = [Paragraph(escape(store_name), styles["Title"]), p(f"Daily orders · {day.isoformat()} · Asia/Kolkata"),
        p(f"{len(orders)} orders | Total INR {sum(o['total'] for o in orders):,} | Non-cancelled INR {sum(o['total'] for o in orders if o['status'] != 'CANCELLED'):,}"), Spacer(1, 20)]
    if not orders:
        story.append(p("No orders for this date."))
    for order in orders:
        heading = [Paragraph(escape(order["order_number"]), styles["Heading2"]), p(f"{local_date(order['created_at'])} | {order['status']}"),
            p(f"Customer: {order['customer_name']} | {order['customer_email']} | {order['customer_phone']}"),
            p(f"Delivery: {order['address']}, {order['city']} - {order['pincode']}"), Spacer(1, 8)]
        story.append(KeepTogether(heading))
        rows = [[p("Product / SKU"), p("Qty"), p("Unit INR"), p("Total INR")]]
        rows += [[p(f"{i['title']} / {i['sku']}"), p(i["quantity"]), p(f"{i['price']:,}"), p(f"{i['quantity'] * i['price']:,}")] for i in order["items"]]
        table = Table(rows, colWidths=[275, 40, 90, 90], repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5f3ef")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 8), ("LINEBELOW", (0, 0), (-1, -1), .5, colors.HexColor("#dfe5e2"))]))
        story += [table, p(f"Order total: INR {order['total']:,}"), Spacer(1, 24)]
    def footer(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(A4[0] - 50, 25, f"Page {document.page} | Confidential order report")
    SimpleDocTemplate(output, pagesize=A4, rightMargin=50, leftMargin=50, topMargin=40, bottomMargin=45).build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()